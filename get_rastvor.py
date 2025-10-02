from openpyxl import load_workbook




def add_rastvor(file_ostatki: str):
    source_file = "rastvor.xlsx"        # откуда копируем строки
    destination_file = file_ostatki  # куда дописываем в конец

    adres = file_ostatki.split("_")[0]
    street = ""

    if adres == 'moscow': #Москва
        street = 'Шолохова'        
    elif adres == 'rostov': #Ростов
        street = 'Островитянова'      
    elif adres == 'kazan': #Казань
        street = 'ул Восстания'      
    elif adres == 'spb': #СПБ
        street = 'пр-кт Обуховской Обороны'      
    elif adres == 'novosib': #Новосибирск
        street = 'ул Гоголя'       
    elif adres == 'ekb': #Екатеринбург
        street = 'ул Норильская'
        

    # открыть книги
    wb_src = load_workbook(source_file, data_only=True)   # data_only=True берёт значения формул
    wb_dst = load_workbook(destination_file)

    # первые листы
    ws_src = wb_src.worksheets[0]
    ws_dst = wb_dst.worksheets[0]

    # сколько столбцов занято в таблице назначения
    dest_cols = ws_dst.max_column

    # если в исходнике есть заголовки в первой строке — пропускаем её
    skip_header = False

    rows_iter = ws_src.iter_rows(values_only=True)
    if skip_header:
        next(rows_iter, None)

    for row in rows_iter:
        # обрезаем до количества столбцов назначения
        trimmed = list(row[:dest_cols])

        # при необходимости дополняем пустыми значениями до dest_cols
        if len(trimmed) < dest_cols:
            trimmed += [None] * (dest_cols - len(trimmed))

        # ⚡ заменяем значение в 4-м столбце (D) на street
        if dest_cols >= 4:
            trimmed[3] = street

        ws_dst.append(trimmed)

    wb_dst.save(destination_file)
    print("Готово: строки дописаны в конец без добавления новых столбцов.")
    print(adres)
    print(street)



